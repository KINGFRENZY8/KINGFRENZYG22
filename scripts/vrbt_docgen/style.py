"""Shared professional styling for Advanta Africa VRBT audit evidence documents."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

NAVY = RGBColor(0x12, 0x2B, 0x4F)
TEAL = RGBColor(0x0E, 0x6B, 0x7A)
GOLD = RGBColor(0xB8, 0x8A, 0x2E)
DARK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x4B, 0x55, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED_CONF = RGBColor(0x8B, 0x1E, 0x1E)

NAVY_HEX = "122B4F"
TEAL_HEX = "0E6B7A"
GOLD_HEX = "B88A2E"
HEADER_ROW_HEX = "122B4F"
ALT_ROW_HEX = "F4F7F8"
BOX_HEX = "E8F2F4"
AMBER_HEX = "FBF4E6"

COMPANY = "Advanta Africa Ltd"
SLOGAN = "Smarter, Faster Messaging"
ADDRESS = "Western Heights, 6th Floor, Karuna Road, Westlands, Nairobi, Kenya"
POSTAL = "P.O. Box 1035-00200, City Square, Nairobi"
PHONE = "+254 111 139 000"
EMAIL = "privacy@advanta.africa"
SUPPORT = "support@advanta.africa"
WEB = "www.advantasms.com | www.advanta.africa"
CEO = "Lawrence Muriuki"
CEO_TITLE = "Chief Executive Officer and Founder"
PROJECT = "Video Ring Back Tune (VRBT) Service"
PARTNER = "Huawei Technologies"
REVIEWER = "TQC / CEG"
DOC_DATE = date(2026, 8, 17)
EFFECTIVE = date(2026, 8, 1)
REVIEW_BY = date(2027, 7, 31)
VERSION = "1.0"


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=DARK, all_caps=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    run.font.all_caps = all_caps
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def shade_cell(cell: _Cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in tcPr.findall(qn("w:shd")):
        tcPr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell: _Cell, color="C5CDD3", sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(child)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def set_cell_margins(cell: _Cell, top=80, bottom=80, left=100, right=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_vertical_align(cell: _Cell, val="center"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), val)
    tcPr.append(vAlign)


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def set_table_widths(table: Table, widths_cm: list[float]):
    table.autofit = False
    table.allow_autofit = False
    total = int(sum(Cm(w) for w in widths_cm))
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    for child in list(tblPr):
        if child.tag == qn("w:tblW"):
            tblPr.remove(child)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for g in list(grid):
            grid.remove(g)
    else:
        grid = OxmlElement("w:tblGrid")
        tblPr.addnext(grid)
    for w in widths_cm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(Cm(w))))
        grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[i])
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            for child in tcPr.findall(qn("w:tcW")):
                tcPr.remove(child)
            tcW = OxmlElement("w:tcW")
            tcW.set(qn("w:w"), str(int(Cm(widths_cm[i]))))
            tcW.set(qn("w:type"), "dxa")
            tcPr.append(tcW)


def set_keep_with_next(paragraph: Paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    kwn = OxmlElement("w:keepNext")
    pPr.append(kwn)


def add_bottom_border(paragraph: Paragraph, color=GOLD_HEX, sz="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_page_number(paragraph: Paragraph):
    run1 = paragraph.add_run("Page ")
    set_run_font(run1, size=8, color=MUTED)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r2 = paragraph.add_run()
    r2._r.append(fld1)
    r2._r.append(instr)
    r2._r.append(fld2)
    set_run_font(r2, size=8, color=MUTED)
    run3 = paragraph.add_run(" of ")
    set_run_font(run3, size=8, color=MUTED)
    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "begin")
    instr2 = OxmlElement("w:instrText")
    instr2.set(qn("xml:space"), "preserve")
    instr2.text = " NUMPAGES "
    fld4 = OxmlElement("w:fldChar")
    fld4.set(qn("w:fldCharType"), "end")
    r4 = paragraph.add_run()
    r4._r.append(fld3)
    r4._r.append(instr2)
    r4._r.append(fld4)
    set_run_font(r4, size=8, color=MUTED)


def new_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.15
    nf = normal._element.get_or_add_rPr()
    rFonts = nf.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rFonts.set(qn("w:eastAsia"), "Calibri")
    return doc


def add_header_footer(doc: Document, doc_no: str, short_title: str):
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.clear()
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run(COMPANY.upper())
    set_run_font(r, size=9, bold=True, color=NAVY)
    r2 = hp.add_run("  |  " + SLOGAN + "  |  ")
    set_run_font(r2, size=8, italic=True, color=MUTED)
    r3 = hp.add_run(doc_no)
    set_run_font(r3, size=8, bold=True, color=TEAL)
    add_bottom_border(hp, GOLD_HEX, "18")
    hp.paragraph_format.space_after = Pt(2)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_bottom_border(fp, NAVY_HEX, "12")
    r = fp.add_run("CONFIDENTIAL  ·  Huawei Supplier Audit Evidence  ·  " + short_title + "  ·  ")
    set_run_font(r, size=8, color=MUTED)
    add_page_number(fp)
    fp.paragraph_format.space_before = Pt(4)


def cover_meta_table(doc: Document, rows: list[tuple[str, str]]):
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [5.5, 11.5])
    for i, (k, v) in enumerate(rows):
        prevent_row_split(table.rows[i])
        c0, c1 = table.rows[i].cells
        shade_cell(c0, NAVY_HEX if i % 2 == 0 else TEAL_HEX)
        shade_cell(c1, "F7F9FA")
        for c in (c0, c1):
            set_cell_borders(c, "FFFFFF", "4")
            set_cell_margins(c, 70, 70, 120, 120)
            set_vertical_align(c)
            c.paragraphs[0].clear()
        p0 = c0.paragraphs[0]
        r = p0.add_run(k)
        set_run_font(r, size=10, bold=True, color=WHITE)
        p1 = c1.paragraphs[0]
        r = p1.add_run(v)
        set_run_font(r, size=10, color=DARK)
    return table


def add_cover(
    doc: Document,
    doc_no: str,
    title: str,
    subtitle: str,
    questionnaire_ref: str,
    evidence_for: str,
):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(COMPANY.upper())
    set_run_font(r, name="Calibri", size=22, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(SLOGAN)
    set_run_font(r, size=12, italic=True, color=TEAL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Bulk SMS  ·  Short Codes  ·  SMS APIs  ·  USSD  ·  WhatsApp APIs  ·  VRBT")
    set_run_font(r, size=9, color=MUTED)

    # Gold rule
    rule = doc.add_table(rows=1, cols=1)
    set_table_widths(rule, [17.0])
    shade_cell(rule.cell(0, 0), GOLD_HEX)
    set_cell_borders(rule.cell(0, 0), GOLD_HEX, "0")
    rule.cell(0, 0).paragraphs[0].paragraph_format.space_after = Pt(0)
    rule.rows[0].height = Cm(0.12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("HUAWEI SUPPLIER AUDIT EVIDENCE PACK")
    set_run_font(r, size=11, bold=True, color=GOLD, all_caps=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    set_run_font(r, size=20, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(subtitle)
    set_run_font(r, size=12, italic=True, color=TEAL)

    banner = doc.add_table(rows=1, cols=1)
    set_table_widths(banner, [17.0])
    shade_cell(banner.cell(0, 0), NAVY_HEX)
    set_cell_borders(banner.cell(0, 0), NAVY_HEX, "0")
    set_cell_margins(banner.cell(0, 0), 90, 90, 140, 140)
    bp = banner.cell(0, 0).paragraphs[0]
    bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = bp.add_run("PROJECT:  " + PROJECT)
    set_run_font(r, size=11, bold=True, color=WHITE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)

    cover_meta_table(
        doc,
        [
            ("Document Number", doc_no),
            ("Questionnaire Reference", questionnaire_ref),
            ("Evidence Addresses", evidence_for),
            ("Document Version", VERSION + "  |  Issued for Audit"),
            ("Document Date", DOC_DATE.strftime("%d %B %Y")),
            ("Effective Date", EFFECTIVE.strftime("%d %B %Y")),
            ("Next Review Date", REVIEW_BY.strftime("%d %B %Y")),
            ("Classification", "Confidential — Restricted to Advanta Africa, Huawei TQC/CEG and authorised auditors"),
            ("Prepared For", f"{PARTNER}  ·  Reviewer: {REVIEWER}  ·  Method: Document Verification"),
            ("Prepared By", "Quality Management Office, Advanta Africa Ltd"),
            ("Approved By", f"{CEO}, {CEO_TITLE}"),
            ("Owner", "Quality Management Responsible Person (QMRP)"),
            ("Related Contract", "Huawei content / VAS cooperation for DIY Audio, Video, Game and E-book scenarios, as applied to VRBT"),
        ],
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(ADDRESS)
    set_run_font(r, size=9, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{POSTAL}  ·  Tel: {PHONE}  ·  {WEB}")
    set_run_font(r, size=9, color=MUTED)

    doc.add_page_break()


def heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    if level == 1:
        add_bottom_border(p, GOLD_HEX, "10")
        r = p.add_run(text)
        set_run_font(r, size=14, bold=True, color=NAVY)
    elif level == 2:
        r = p.add_run(text)
        set_run_font(r, size=12, bold=True, color=TEAL)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, bold=True, italic=True, color=NAVY)
    return p


def body(doc: Document, text: str, *, bold=False, italic=False, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    set_run_font(r, size=11, bold=bold, italic=italic, color=DARK)
    return p


def bullet(doc: Document, text: str, level: int = 0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0 + level * 0.6)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("•  " + text)
    set_run_font(r, size=11, color=DARK)
    return p


def numbered(doc: Document, n: int, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{n}.  {text}")
    set_run_font(r, size=11, color=DARK)
    return p


def callout(doc: Document, title: str, text: str, fill=BOX_HEX):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [17.0])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_borders(cell, TEAL_HEX, "8")
    set_cell_margins(cell, 100, 100, 140, 140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title.upper())
    set_run_font(r, size=9, bold=True, color=TEAL)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p2.add_run(text)
    set_run_font(r, size=10, color=DARK)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(8)
    return table


def styled_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    if widths is None:
        col_w = round(17.0 / len(headers), 2)
        widths = [col_w] * len(headers)
        widths[-1] = round(17.0 - sum(widths[:-1]), 2)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, widths)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, HEADER_ROW_HEX)
        set_cell_borders(cell, HEADER_ROW_HEX, "4")
        set_cell_margins(cell, 60, 60, 80, 80)
        set_vertical_align(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_run_font(r, size=9, bold=True, color=WHITE)
    prevent_row_split(table.rows[0])
    for ri, row in enumerate(rows):
        prevent_row_split(table.rows[ri + 1])
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            shade_cell(cell, ALT_ROW_HEX if ri % 2 else "FFFFFF")
            set_cell_borders(cell, "D0D7DE", "4")
            set_cell_margins(cell, 55, 55, 80, 80)
            set_vertical_align(cell, "top")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val))
            set_run_font(r, size=9, color=DARK)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(8)
    return table


def kv_table(doc: Document, pairs: list[tuple[str, str]], widths=None):
    if widths is None:
        widths = [5.2, 11.8]
    table = doc.add_table(rows=len(pairs), cols=2)
    set_table_widths(table, widths)
    for i, (k, v) in enumerate(pairs):
        prevent_row_split(table.rows[i])
        c0, c1 = table.rows[i].cells
        shade_cell(c0, NAVY_HEX)
        shade_cell(c1, ALT_ROW_HEX if i % 2 else "FFFFFF")
        set_cell_borders(c0, NAVY_HEX, "4")
        set_cell_borders(c1, "D0D7DE", "4")
        set_cell_margins(c0, 55, 55, 90, 90)
        set_cell_margins(c1, 55, 55, 90, 90)
        set_vertical_align(c0, "top")
        set_vertical_align(c1, "top")
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r = p0.add_run(k)
        set_run_font(r, size=9, bold=True, color=WHITE)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r = p1.add_run(v)
        set_run_font(r, size=9, color=DARK)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(8)
    return table


def add_picture_centered(doc: Document, path: str | Path, width_cm=16.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    return p


def signature_block(doc: Document, signatories: list[tuple[str, str, str]]):
    heading(doc, "Document Approval", 1)
    body(
        doc,
        "This document is issued under the authority of Advanta Africa Ltd for the VRBT programme. "
        "Wet-ink or authorised electronic signatures below confirm that the content is accurate, "
        "that the described controls are in operation, and that the document may be presented to Huawei TQC/CEG.",
    )
    cols = min(3, len(signatories))
    table = doc.add_table(rows=1, cols=cols)
    width = round(17.0 / cols, 2)
    set_table_widths(table, [width] * cols)
    for i, (role, name, date_text) in enumerate(signatories[:cols]):
        cell = table.rows[0].cells[i]
        set_cell_borders(cell, "C5CDD3", "8")
        set_cell_margins(cell, 120, 120, 120, 120)
        shade_cell(cell, "FFFFFF")
        p = cell.paragraphs[0]
        r = p.add_run(role.upper())
        set_run_font(r, size=8, bold=True, color=TEAL)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(22)
        r = p2.add_run("Signature: ____________________")
        set_run_font(r, size=9, color=DARK)
        p3 = cell.add_paragraph()
        r = p3.add_run("Name: " + name)
        set_run_font(r, size=9, color=DARK)
        p4 = cell.add_paragraph()
        r = p4.add_run("Date: " + date_text)
        set_run_font(r, size=9, color=DARK)
    if len(signatories) > cols:
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)
        table2 = doc.add_table(rows=1, cols=len(signatories) - cols)
        n = len(signatories) - cols
        w = round(17.0 / n, 2)
        set_table_widths(table2, [w] * n)
        for i, (role, name, date_text) in enumerate(signatories[cols:]):
            cell = table2.rows[0].cells[i]
            set_cell_borders(cell, "C5CDD3", "8")
            set_cell_margins(cell, 120, 120, 120, 120)
            p = cell.paragraphs[0]
            r = p.add_run(role.upper())
            set_run_font(r, size=8, bold=True, color=TEAL)
            p2 = cell.add_paragraph()
            p2.paragraph_format.space_before = Pt(22)
            r = p2.add_run("Signature: ____________________")
            set_run_font(r, size=9, color=DARK)
            p3 = cell.add_paragraph()
            r = p3.add_run("Name: " + name)
            set_run_font(r, size=9, color=DARK)
            p4 = cell.add_paragraph()
            r = p4.add_run("Date: " + date_text)
            set_run_font(r, size=9, color=DARK)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(8)


def revision_history(doc: Document, extra_rows=None):
    heading(doc, "Document Control and Revision History", 1)
    rows = [
        ["0.1", "15 Jul 2026", "Quality Office", "Draft framework for VRBT Huawei evidence pack"],
        ["0.9", "05 Aug 2026", "QMRP / DPO / Content Compliance", "Internal review against questionnaire key review points"],
        ["1.0", "17 Aug 2026", "CEO", "Issued for Huawei TQC/CEG document verification"],
    ]
    if extra_rows:
        rows.extend(extra_rows)
    styled_table(doc, ["Version", "Date", "Author / Authority", "Description"], rows, [2.2, 3.0, 5.0, 6.8])


def distribution(doc: Document):
    heading(doc, "Distribution List", 2)
    styled_table(
        doc,
        ["Copy", "Recipient", "Purpose"],
        [
            ["Master", "Quality Management Office (controlled repository)", "Document of record"],
            ["1", "Huawei TQC / CEG auditor", "Supplier document verification"],
            ["2", "Advanta Africa CEO / Management", "Accountability and sponsorship"],
            ["3", "Data Protection Officer", "Privacy and cross-border controls"],
            ["4", "Content Compliance Officer", "Content, IP, ethics and go-live review"],
            ["5", "VRBT Product and Operations", "Day-to-day implementation"],
        ],
        [2.2, 7.5, 7.3],
    )


def start_doc(doc_no: str, title: str, subtitle: str, qref: str, evidence_for: str, short: str) -> Document:
    doc = new_document()
    add_header_footer(doc, doc_no, short)
    add_cover(doc, doc_no, title, subtitle, qref, evidence_for)
    return doc


def save_doc(doc: Document, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
